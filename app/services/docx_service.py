from docx import Document
import logging
import os

logger = logging.getLogger(__name__)


def extract_docx(file_path: str) -> str:
    """
    Extracts text from a DOCX file, ensuring that both standard paragraphs
    and embedded tables (crucial for RFQs and BOQs) are captured.
    """
    if not os.path.exists(file_path):
        # 🔧 FIX: raise instead of returning "" silently — a missing file
        # was previously indistinguishable downstream from a file that
        # parsed but contained no usable text, which produces a vague
        # "No meaningful text extracted" error instead of the real cause.
        logger.error(f"❌ DOCX not found at path: {file_path}")
        raise FileNotFoundError(f"DOCX file not found: {file_path}")

    try:
        doc = Document(file_path)
        content_chunks = []

        # 1️⃣ Extract standard paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                content_chunks.append(para.text.strip())

        # 2️⃣ Extract tables (CRITICAL FOR RFQs)
        # Note: doc.tables only returns top-level tables — a table nested
        # inside a table cell (uncommon, but seen in some RFQ templates)
        # won't be picked up here.
        if doc.tables:
            content_chunks.append("\n--- DOCUMENT TABLES ---")
            for table_idx, table in enumerate(doc.tables):
                content_chunks.append(f"\n[Table {table_idx + 1}]")

                for row in table.rows:
                    # NOTE: python-docx repeats cell.text for every cell in
                    # a merged region (it doesn't collapse merges), so a
                    # merged header can appear duplicated in row_data below.
                    # Deduplicating consecutive identical cells here would
                    # risk collapsing legitimately repeated values (e.g. two
                    # columns that both happen to say "N/A"), so this is
                    # left as a known limitation rather than "fixed" blindly.
                    row_data = [cell.text.replace('\n', ' ').strip() for cell in row.cells]

                    if any(row_data):
                        content_chunks.append(" | ".join(row_data))

        logger.info(f"✅ Successfully extracted text and {len(doc.tables)} tables from DOCX.")
        return "\n".join(content_chunks)

    except Exception as e:
        # 🔧 FIX: exc_info=True so the actual traceback lands in logs,
        # not just the message — matters here since python-docx failures
        # (corrupt file, unsupported format variant, etc.) are common
        # enough to need more than a one-line error to debug.
        logger.error(f"❌ Critical failure extracting DOCX {file_path}. Error: {e}", exc_info=True)
        raise ValueError(f"Failed to read DOCX file: {e}") from e
